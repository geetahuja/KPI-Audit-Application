"""Read-only discovery. Never edits source workbooks. openpyxl is not a calculator."""
import argparse, json, re, pathlib, datetime, warnings
import openpyxl
from openpyxl.formula.translate import Translator
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

def role(name):
    n=name.lower()
    for term, value in [('data add','input'),('dashboard','dashboard'),('benchmark','benchmark'),('evidence','evidence'),('graph','graph'),('formulae','configuration'),('company','configuration'),('responsibility','management'),('activities','management'),('index sheet','input')]:
        if term in n:return value
    return 'calculation'

def extract(path):
    w=openpyxl.load_workbook(path, data_only=False)
    v=openpyxl.load_workbook(path, data_only=True)
    out={'name':path.name,'role':'Master' if any('Productivity (' in s for s in w.sheetnames) else 'MCLS','sheets':[], 'namedRanges':list(w.defined_names),'externalLinks':len(w._external_links)}
    for s in w:
        cells={}
        for row in s:
            for c in row:
                if c.value is None:continue
                val=c.value
                f=(val if isinstance(val,str) else getattr(val,'text','')) if c.data_type=='f' else None
                raw=v[s.title][c.coordinate].value if f else val
                if isinstance(raw,(datetime.datetime,datetime.date)):raw=raw.isoformat()
                cells[c.coordinate]={'v':raw,'f':f,'t':c.data_type,'z':c.number_format}
        out['sheets'].append({'name':s.title,'role':role(s.title),'state':s.sheet_state,'rows':s.max_row,'cols':s.max_column,'merges':[str(x) for x in s.merged_cells.ranges],'validations':len(s.data_validations.dataValidation),'tables':list(s.tables),'charts':len(s._charts),'conditionalFormats':len(s.conditional_formatting),'cells':cells})
    return out

def main():
    p=argparse.ArgumentParser();p.add_argument('folder');p.add_argument('output');a=p.parse_args()
    output=pathlib.Path(a.output);output.mkdir(parents=True,exist_ok=True)
    inventory=[]; reference=[]
    for path in sorted(pathlib.Path(a.folder).glob('*.xlsx')):
        w=extract(path); (output/(path.stem+'.json')).write_text(json.dumps(w,default=str))
        sheets=[]
        for s in w['sheets']:
            formulas=[(k,c['f']) for k,c in s['cells'].items() if c['f']]
            technical=[k for k,f in formulas if re.search('DUMMYFUNCTION|IMPORTRANGE',f,re.I)]
            deps=sorted(set(x for _,f in formulas for x in re.findall(r"'([^']+)'!",f)))
            sheets.append({k:v for k,v in s.items() if k!='cells'}|{'formulaCount':len(formulas),'conversionArtifacts':len(technical),'dependencies':deps})
            if path.name.startswith('Entech'):
                for row in range(3,min(s['rows'],65)+1):
                    label=s['cells'].get('C'+str(row),{}).get('v')
                    unit=s['cells'].get('B'+str(row),{}).get('v')
                    examples=[{'cell':c,'formula':f} for c,f in formulas if re.fullmatch(r'[D-F]'+str(row),c) and not re.search('DUMMYFUNCTION|IMPORTRANGE',f,re.I)]
                    if s['role']=='calculation' and isinstance(label,str) and examples:
                        reference.append({'sheet':s['name'],'label':label.strip(),'unit':unit,'examples':examples})
        inventory.append({'name':w['name'],'role':w['role'],'sheets':sheets,'namedRanges':w['namedRanges'],'externalLinks':w['externalLinks']})
        print(path.name,len(sheets),'sheets',sum(s['formulaCount'] for s in sheets),'formulas',flush=True)
    (output/'inventory.json').write_text(json.dumps(inventory,indent=2))
    (output/'reference-patterns.json').write_text(json.dumps(reference,indent=2))

if __name__=='__main__':main()
